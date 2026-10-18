# -*- coding: utf-8 -*-
"""
只读诊断：为什么"逐赛事加和"结构对 HLTV 官方 7-20 名中段次序复现有限。

不修改任何现有文件。读取 if_sim_2025_events.xlsx / baseline，复现模拟器 v1.2.0
算法，然后分解误差来源。
"""
import os
import sys
import json
from collections import defaultdict

import openpyxl
import numpy as np

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 项目根 = .analysis/ 的上一级
DATA = os.path.join(BASE, "data")

# ---------- 模拟器 v1.2.0 默认权重（来自 config/defaults.ts） ----------
EVENT_LEVEL_WEIGHTS = {
    "Major": 1.0, "超级精英赛": 0.8, "精英赛": 0.75,
    "大型赛事": 0.65, "中型赛事": 0.5, "次级赛事": 0.45, "不纳入": 0.0,
}
PLACEMENT_ADJUST = {"冠军": 2, "亚军": 0.5, "四强": 0, "八强": -0.5, "小组淘汰": -1}
HONOR_BASE = {"MVP": 10, "1": 8, "2": 6, "3": 4}
RATING_TIERS = [(1.3, 2), (1.2, 1), (1.1, 0.5)]
RATING_BELOW = [(1.05, -0.5), (1.0, -1)]
SCORE_MIN, SCORE_MAX = -2, 11
MIN_MAPS, MIN_EVENT_MAPS = 20, 3


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def load_rows(path, sheet=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if r is None or all(v is None for v in r):
            continue
        out.append(dict(zip(header, r)))
    return out


def rating_bonus(rating):
    bonus = 0.0
    for th, b in RATING_TIERS:
        if rating >= th:
            bonus = max(bonus, b)
    for th, b in RATING_BELOW:
        if rating < th:
            bonus = min(bonus, b)
    return bonus


def event_score(row):
    honor = str(row.get("荣誉类型") or "无").strip()
    rank = int(num(row.get("荣誉顺位")))
    if honor == "MVP":
        base = HONOR_BASE["MVP"]
    elif honor == "EVP":
        base = HONOR_BASE.get(str(rank), 0)
    else:
        base = 0
    place = str(row.get("队伍名次") or "").strip()
    s = base + rating_bonus(num(row.get("Rating"))) + PLACEMENT_ADJUST.get(place, 0)
    return min(SCORE_MAX, max(SCORE_MIN, s))


def compute_sim(rows):
    totals = defaultdict(float)
    maps_sum = defaultdict(float)
    detail = defaultdict(list)
    for r in rows:
        # 用选手名做键：baseline 只提供名字，ID 在赛事表里是数字、无法直接对齐
        pid = str(r.get("选手名") or r.get("选手ID") or "").strip()
        if not pid:
            continue
        maps = num(r.get("地图数"))
        if maps and maps < MIN_EVENT_MAPS:
            continue
        maps_sum[pid] += maps
        w = num(r.get("赛事权重")) or EVENT_LEVEL_WEIGHTS.get(
            str(r.get("赛事级别") or "").strip(), 0)
        if w == 0:
            continue
        s = event_score(r)
        totals[pid] += s * w
        detail[pid].append((str(r.get("赛事名称") or ""), s, w, s * w))
    out = {}
    for pid, tot in totals.items():
        if maps_sum[pid] and maps_sum[pid] < MIN_MAPS:
            continue
        out[pid] = (tot, maps_sum[pid], detail[pid])
    return out


def spearman(a, b):
    a, b = np.asarray(a, float), np.asarray(b, float)
    if len(a) < 2:
        return float("nan")
    ra = np.argsort(np.argsort(a)).astype(float)
    rb = np.argsort(np.argsort(b)).astype(float)
    ra -= ra.mean(); rb -= rb.mean()
    d = np.sqrt((ra ** 2).sum() * (rb ** 2).sum())
    return float((ra * rb).sum() / d) if d else float("nan")


def main():
    ev_path = os.path.join(DATA, "if_sim_2025_events.xlsx")
    bl_path = os.path.join(DATA, "if_sim_2025_baseline.xlsx")
    rows = load_rows(ev_path)
    baseline = load_rows(bl_path)
    print(f"赛事行数={len(rows)}  基准行数={len(baseline)}")
    print(f"赛事数据列: {list(rows[0].keys())}")
    print(f"基准列: {list(baseline[0].keys())}")

    sim = compute_sim(rows)
    official = {}
    for r in baseline:
        name = str(r.get("选手名") or "").strip()
        rk = int(num(r.get("名次")))
        if name and rk:
            official[name] = rk

    ranked = sorted(sim.items(), key=lambda kv: -kv[1][0])
    pred_rank = {pid: i + 1 for i, (pid, _) in enumerate(ranked)}
    print(f"\n模拟器产出选手数={len(ranked)}  官方Top20数={len(official)}")

    print("\n=== 官方 Top20 vs 模拟器名次（逐名） ===")
    print(f"{'官方名次':>8} {'选手':<16} {'模拟名次':>8} {'模拟总分':>9} {'偏差':>6}")
    devs = []
    for name, rk in sorted(official.items(), key=lambda kv: kv[1]):
        pr = pred_rank.get(name)
        tot = sim.get(name, (None,))[0]
        if pr is None:
            print(f"{rk:>8} {name:<16} {'未进榜':>8} {'-':>9} {'-':>6}")
            devs.append((rk, None))
        else:
            d = pr - rk
            devs.append((rk, d))
            print(f"{rk:>8} {name:<16} {pr:>8} {tot:>9.2f} {d:>+6}")

    # 分档统计误差
    print("\n=== 分档误差（官方名次区间） ===")
    bands = [(1, 6, "头部 1-6"), (7, 13, "中段 7-13"), (14, 20, "后段 14-20")]
    for lo, hi, label in bands:
        ds = [d for rk, d in devs if lo <= rk <= hi and d is not None]
        miss = sum(1 for rk, d in devs if lo <= rk <= hi and d is None)
        if ds:
            ad = np.abs(ds)
            print(f"{label:<12} n={len(ds):2d} 平均绝对偏差={ad.mean():5.2f} "
                  f"最大={ad.max():5.2f} 未进榜={miss} 平均有符号偏差={np.mean(ds):+5.2f}")
        else:
            print(f"{label:<12} 无数据 未进榜={miss}")

    # 头部命中率
    print("\n=== Top-N 重叠 ===")
    for n in (5, 10, 20):
        off_top = {nm for nm, rk in official.items() if rk <= n}
        pred_top = {pid for pid, _ in ranked[:n]}
        print(f"Top{n:>2}: 重叠 {len(off_top & pred_top)}/{n}")

    # 整体 Spearman
    common = [nm for nm in official if nm in pred_rank]
    if len(common) >= 2:
        sp = spearman([official[nm] for nm in common], [pred_rank[nm] for nm in common])
        print(f"\n20 人共同集合 Spearman = {sp:.3f}")

    # 谁被挤进来了：模拟器前20里没有官方名次的
    print("\n=== 模拟器 Top20 中不在官方 Top20 的选手 ===")
    for pid, (tot, mp, det) in ranked[:20]:
        if pid not in official:
            print(f"  模拟#{pred_rank[pid]:<3} {pid:<16} 总分={tot:6.2f} 图数={mp:.0f}")

    # 关键：中段选手的分数差有多小
    print("\n=== 官方 7-20 名在模拟器中的分数分布（分数差 = 区分度） ===")
    seg = [(nm, sim[nm][0], official[nm], pred_rank[nm])
           for nm in official if 7 <= official[nm] <= 20 and nm in sim]
    seg.sort(key=lambda x: -x[1])
    if seg:
        scores = [s for _, s, _, _ in seg]
        print(f"  人数={len(seg)} 最高={max(scores):.2f} 最低={min(scores):.2f} "
              f"极差={max(scores)-min(scores):.2f}")
        gaps = [scores[i] - scores[i + 1] for i in range(len(scores) - 1)]
        if gaps:
            print(f"  相邻名次平均分差={np.mean(gaps):.3f} 最小分差={min(gaps):.3f}")
        for nm, s, ork, prk in seg:
            print(f"    官方#{ork:<3} {nm:<16} 模拟#{prk:<3} 总分={s:6.2f}")

    # 统计分数并列/极近的情况
    all_scores = sorted([v[0] for v in sim.values()], reverse=True)[:40]
    tied = 0
    for i in range(len(all_scores) - 1):
        if abs(all_scores[i] - all_scores[i + 1]) < 0.01:
            tied += 1
    print(f"\n前40名中相邻分差<0.01 的对数 = {tied}")


if __name__ == "__main__":
    main()
