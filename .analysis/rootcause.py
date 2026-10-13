# -*- coding: utf-8 -*-
"""
只读诊断 2：中段误差的根因分解。

假设 A：逐赛事加和放大了"荣誉事件数"，官方是 7-20 名看的是全年聚合质量而非荣誉计数。
假设 B：中段选手分数被"冠军队伍的名次分"抬高（冠军队多名选手共享 +2/赛事）。
假设 C：缺少"TOP5/10/20 对手强度"维度——官方中段排序主要靠对手质量与稳定性。
假设 D：rating 分档过于粗糙（只有 3 档正分），中段 1.10-1.25 的密集区间被压平。

用真实数据逐条量化。
"""
import os
from collections import defaultdict
import openpyxl
import numpy as np

BASE = r"D:\demo\iftv"
DATA = os.path.join(BASE, "data")

EVENT_LEVEL_WEIGHTS = {
    "Major": 1.0, "超级精英赛": 0.8, "精英赛": 0.75,
    "大型赛事": 0.65, "中型赛事": 0.5, "次级赛事": 0.45, "不纳入": 0.0,
}
PLACEMENT_ADJUST = {"冠军": 2, "亚军": 0.5, "四强": 0, "八强": -0.5, "小组淘汰": -1}
HONOR_BASE = {"MVP": 10, "1": 8, "2": 6, "3": 4}
RATING_TIERS = [(1.3, 2), (1.2, 1), (1.1, 0.5)]
RATING_BELOW = [(1.05, -0.5), (1.0, -1)]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def load_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if r and any(v is not None for v in r)]


def rating_bonus(r):
    b = 0.0
    for th, bo in RATING_TIERS:
        if r >= th:
            b = max(b, bo)
    for th, bo in RATING_BELOW:
        if r < th:
            b = min(b, bo)
    return b


def decompose(rows):
    """把每位选手的总分拆成 荣誉分/rating分/名次分 三条通道（含加权）"""
    ch = defaultdict(lambda: {"honor": 0.0, "rating": 0.0, "place": 0.0,
                              "n_events": 0, "n_mvp": 0, "n_evp": 0,
                              "maps": 0.0, "ratings": [], "honor_events": []})
    for r in rows:
        nm = str(r.get("选手名") or "").strip()
        if not nm:
            continue
        maps = num(r.get("地图数"))
        if maps and maps < 3:
            continue
        w = num(r.get("赛事权重")) or EVENT_LEVEL_WEIGHTS.get(
            str(r.get("赛事级别") or "").strip(), 0)
        if w == 0:
            continue
        c = ch[nm]
        c["maps"] += maps
        c["n_events"] += 1
        rat = num(r.get("Rating"))
        c["ratings"].append(rat)
        honor = str(r.get("荣誉类型") or "无").strip()
        rank = int(num(r.get("荣誉顺位")))
        hb = HONOR_BASE["MVP"] if honor == "MVP" else (
            HONOR_BASE.get(str(rank), 0) if honor == "EVP" else 0)
        if honor == "MVP":
            c["n_mvp"] += 1
            c["honor_events"].append(str(r.get("赛事名称")))
        elif honor == "EVP":
            c["n_evp"] += 1
        place = str(r.get("队伍名次") or "").strip()
        c["honor"] += hb * w
        c["rating"] += rating_bonus(rat) * w
        c["place"] += PLACEMENT_ADJUST.get(place, 0) * w
    for nm, c in ch.items():
        c["total"] = c["honor"] + c["rating"] + c["place"]
        c["avg_rating"] = float(np.mean(c["ratings"])) if c["ratings"] else 0.0
    return ch


def main():
    rows = load_rows(os.path.join(DATA, "if_sim_2025_events.xlsx"))
    baseline = load_rows(os.path.join(DATA, "if_sim_2025_baseline.xlsx"))
    official = {str(r.get("选手名")).strip(): int(num(r.get("名次"))) for r in baseline}

    ch = decompose(rows)
    ranked = sorted(ch.items(), key=lambda kv: -kv[1]["total"])
    pred = {nm: i + 1 for i, (nm, _) in enumerate(ranked)}

    print("=" * 100)
    print("官方 Top20 的三通道分数分解（按官方名次排序）")
    print("=" * 100)
    print(f"{'官方':>4} {'选手':<14} {'模拟':>4} {'总分':>7} {'荣誉':>7} {'rating':>7} "
          f"{'名次':>7} {'MVP':>4} {'EVP':>4} {'场次':>4} {'均Rating':>8} {'图数':>5}")
    for nm, rk in sorted(official.items(), key=lambda kv: kv[1]):
        c = ch.get(nm)
        if not c:
            print(f"{rk:>4} {nm:<14} 无数据")
            continue
        print(f"{rk:>4} {nm:<14} {pred[nm]:>4} {c['total']:>7.2f} {c['honor']:>7.2f} "
              f"{c['rating']:>7.2f} {c['place']:>7.2f} {c['n_mvp']:>4} {c['n_evp']:>4} "
              f"{c['n_events']:>4} {c['avg_rating']:>8.3f} {c['maps']:>5.0f}")

    # ---- 假设 A：荣誉事件数在中段是否有区分力？ ----
    print("\n" + "=" * 100)
    print("假设 A：中段（官方7-20）误差 vs 荣誉通道占比")
    print("=" * 100)
    mid = [(nm, ch[nm], official[nm], pred[nm]) for nm in official
           if 7 <= official[nm] <= 20 and nm in ch]
    mid.sort(key=lambda x: -abs(x[3] - x[2]))
    for nm, c, ork, prk in mid:
        share = c["honor"] / c["total"] * 100 if c["total"] else 0
        print(f"  官方#{ork:<3} {nm:<14} 偏差={prk-ork:+3d}  荣誉占比={share:5.1f}%  "
              f"荣誉分={c['honor']:6.2f}  MVP={c['n_mvp']} EVP={c['n_evp']}")
    big_err = [abs(prk - ork) for nm, c, ork, prk in mid if abs(prk - ork) >= 5]
    sml_err = [abs(prk - ork) for nm, c, ork, prk in mid if abs(prk - ork) < 5]
    print(f"\n  大误差(>=5)人数={len(big_err)}  小误差(<5)人数={len(sml_err)}")

    # ---- 假设 B：名次分是否在抬高低 rating 选手？ ----
    print("\n" + "=" * 100)
    print("假设 B：名次分通道 vs 平均 rating（冠军队多人共享名次分）")
    print("=" * 100)
    print(f"{'选手':<14} {'场均Rating':>10} {'名次分':>8} {'图数':>6} {'官方':>5} {'模拟':>5}")
    for nm, c in sorted(ch.items(), key=lambda kv: -kv[1]["place"])[:15]:
        o = official.get(nm, "-")
        print(f"{nm:<14} {c['avg_rating']:>10.3f} {c['place']:>8.2f} "
              f"{c['maps']:>6.0f} {str(o):>5} {pred[nm]:>5}")

    # ---- 假设 D：rating 分档的区分度 ----
    print("\n" + "=" * 100)
    print("假设 D：中段选手平均 rating 分布（分档是否压平了差异）")
    print("=" * 100)
    mids = [(nm, ch[nm]["avg_rating"], official[nm]) for nm in official
            if 7 <= official[nm] <= 20 and nm in ch]
    mids.sort(key=lambda x: -x[1])
    for nm, ar, ork in mids:
        print(f"  官方#{ork:<3} {nm:<14} 均Rating={ar:.3f}  → 档位加分={rating_bonus(ar):+.1f}")
    ars = [ar for _, ar, _ in mids]
    print(f"\n  中段均Rating: 最高={max(ars):.3f} 最低={min(ars):.3f} 极差={max(ars)-min(ars):.3f}")
    tiers_hit = [rating_bonus(a) for a in ars]
    print(f"  落入档位: {sorted(set(tiers_hit))}  → 档位只区分出 {len(set(tiers_hit))} 种取值")

    # 官方名次 vs 各通道的 Spearman（在 7-20 段内）
    print("\n" + "=" * 100)
    print("中段（官方7-20）内，各通道与官方名次的相关性")
    print("=" * 100)
    seg = [(nm, ch[nm], official[nm]) for nm in official if 7 <= official[nm] <= 20 and nm in ch]
    oranks = np.array([o for _, _, o in seg], float)

    def sp(vals):
        v = np.array(vals, float)
        rv = np.argsort(np.argsort(v)).astype(float)
        ro = np.argsort(np.argsort(oranks)).astype(float)
        rv -= rv.mean(); ro -= ro.mean()
        den = np.sqrt((rv ** 2).sum() * (ro ** 2).sum())
        return float((rv * ro).sum() / den) if den else float("nan")

    for label, vals in [
        ("总分(当前算法)", [c["total"] for _, c, _ in seg]),
        ("荣誉通道", [c["honor"] for _, c, _ in seg]),
        ("rating通道", [c["rating"] for _, c, _ in seg]),
        ("名次通道", [c["place"] for _, c, _ in seg]),
        ("平均Rating", [c["avg_rating"] for _, c, _ in seg]),
        ("参赛场次", [float(c["n_events"]) for _, c, _ in seg]),
        ("荣誉事件数(MVP+EVP)", [float(c["n_mvp"] + c["n_evp"]) for _, c, _ in seg]),
    ]:
        print(f"  {label:<22} Spearman={sp(vals):+.3f}")


if __name__ == "__main__":
    main()
