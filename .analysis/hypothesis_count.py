# -*- coding: utf-8 -*-
"""
只读诊断 5：验证核心结构假设 —— "逐赛事加和"的最优解是"荣誉事件计数"结构。

观察（structure.py）：
  中段 Spearman 在 {荣誉+名次} 结构下 = +0.209，在 {rating+名次} 下 = +0.319，
  但两者叠加反而降到 +0.231。说明两通道互相干扰，而 rating 通道的
  "信息" 被荣誉通道的 "量级" 压过。

核心假设：逐赛事加和的排序 ≫ 由"荣誉事件数"决定，因为
  荣誉基础分 (4~10) 远大于单赛事 rating 分档 (0~2) 与名次分 (-1~+2)。
  一次 MVP(10分) ≈ 5 次 rating 顶档(2分) ≈ 全年 20+ 场的 rating 差异。

检验方式：直接比较"荣誉事件数"作为唯一排序依据 vs 现行算法的中段表现。
若"数 MVP/EVP 个数"就能达到相近甚至更好的中段 Spearman，
则说明现行算法在中段本质上就是在数荣誉，只是数得不够好。
"""
import os
from collections import defaultdict
import openpyxl
import numpy as np

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 项目根 = .analysis/ 的上一级
DATA = os.path.join(BASE, "data")
EVENT_LEVEL_WEIGHTS = {
    "Major": 1.0, "超级精英赛": 0.8, "精英赛": 0.75,
    "大型赛事": 0.65, "中型赛事": 0.5, "次级赛事": 0.45, "不纳入": 0.0,
}
PLACEMENT_ADJUST = {"冠军": 2, "亚军": 0.5, "四强": 0, "八强": -0.5, "小组淘汰": -1}
RATING_TIERS = [(1.3, 2), (1.2, 1), (1.1, 0.5)]
RATING_BELOW = [(1.05, -0.5), (1.0, -1)]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def load_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if r and any(v is not None for v in r)]


def rating_bonus(r):
    b = 0.0
    for th, bo in RATING_TIERS:
        if r >= th:
            b = max(b, bo)
    for th, bo in RATING_BELOW:
        if r < th:
            b = min(b, bo)
    return b


def spearman(a, b):
    a, b = np.asarray(a, float), np.asarray(b, float)
    ra = np.argsort(np.argsort(a)).astype(float)
    rb = np.argsort(np.argsort(b)).astype(float)
    ra -= ra.mean(); rb -= rb.mean()
    d = np.sqrt((ra ** 2).sum() * (rb ** 2).sum())
    return float((ra * rb).sum() / d) if d else float("nan")


def analyze(ev_path, bl_path, year):
    rows = load_rows(ev_path)
    official = {str(r.get("选手名")).strip(): int(num(r.get("名次")))
                for r in load_rows(bl_path)}
    acc = defaultdict(lambda: defaultdict(float))
    ratings = defaultdict(list)
    for r in rows:
        nm = str(r.get("选手名") or "").strip()
        if not nm:
            continue
        maps = num(r.get("地图数"))
        if maps and maps < 3:
            continue
        w = num(r.get("赛事权重")) or EVENT_LEVEL_WEIGHTS.get(
            str(r.get("赛事级别") or "").strip(), 0)
        if w == 0:
            continue
        a = acc[nm]
        a["maps"] += maps
        ratings[nm].append(num(r.get("Rating")))
        honor = str(r.get("荣誉类型") or "无").strip()
        rank = int(num(r.get("荣誉顺位")))
        hb = 10 if honor == "MVP" else ({1: 8, 2: 6, 3: 4}.get(rank, 0) if honor == "EVP" else 0)
        a["honor"] += hb * w
        a["rating"] += rating_bonus(num(r.get("Rating"))) * w
        a["place"] += PLACEMENT_ADJUST.get(str(r.get("队伍名次") or "").strip(), 0) * w
        if honor == "MVP":
            a["n_mvp"] += 1
        elif honor == "EVP":
            a["n_evp"] += 1

    base = {nm: a["honor"] + a["rating"] + a["place"]
            for nm, a in acc.items() if a["maps"] >= 20}
    avg = {nm: float(np.mean(ratings[nm])) for nm in base}

    def report(label, scores):
        rk = sorted(scores.items(), key=lambda kv: -kv[1])
        pred = {nm: i + 1 for i, (nm, _) in enumerate(rk)}
        mid = [nm for nm in official if 7 <= official[nm] <= 20 and nm in pred]
        head = [nm for nm in official if official[nm] <= 6 and nm in pred]
        top20 = len({nm for nm in official if nm in list(pred)[:20]})
        if not mid:
            print(f"  {label:<38} 无中段数据")
            return
        mae = float(np.mean([abs(pred[nm] - official[nm]) for nm in mid]))
        sp = spearman([official[nm] for nm in mid], [pred[nm] for nm in mid])
        hmae = float(np.mean([abs(pred[nm] - official[nm]) for nm in head])) if head else float("nan")
        print(f"  {label:<38} Top20={top20:>2}/20 头部MAE={hmae:4.2f} "
              f"中段MAE={mae:5.2f} 中段sp={sp:+.3f}")

    print("=" * 104)
    print(f"{year} 年：不同排序依据的中段表现对比")
    print("=" * 104)
    report("现行 v1.2.0（荣誉+rating档+名次）", base)
    report("仅数荣誉事件数 (MVP+EVP 个数)", {nm: float(acc[nm]["n_mvp"] + acc[nm]["n_evp"]) for nm in base})
    report("仅数 MVP 个数", {nm: acc[nm]["n_mvp"] for nm in base})
    report("荣誉加权(赛事级别)计数", {nm: acc[nm]["honor"] for nm in base})
    report("全年平均 Rating（无任何加权）", avg)
    report("平均Rating + 名次通道", {nm: avg[nm] * 40 + acc[nm]["place"] for nm in base})
    report("平均Rating + 荣誉 + 名次", {nm: avg[nm] * 40 + base[nm] for nm in base})
    report("平均Rating×40（连续，无权名次）", {nm: avg[nm] * 40 for nm in base})

    # 相关性矩阵
    print(f"\n  各量与官方名次（中段7-20）的 Spearman：")
    mid = [nm for nm in official if 7 <= official[nm] <= 20 and nm in base]
    oranks = [official[nm] for nm in mid]
    for label, vals in [
        ("现行总分", [base[nm] for nm in mid]),
        ("荣誉事件数", [acc[nm]["n_mvp"] + acc[nm]["n_evp"] for nm in mid]),
        ("荣誉加权分", [acc[nm]["honor"] for nm in mid]),
        ("平均Rating", [avg[nm] for nm in mid]),
        ("rating档位分", [acc[nm]["rating"] for nm in mid]),
        ("名次分", [acc[nm]["place"] for nm in mid]),
        ("参赛场次", [acc[nm]["maps"] for nm in mid]),
    ]:
        # 注意：官方名次越小越好，故取负相关为"方向正确"
        sp = spearman(oranks, vals)
        print(f"    {label:<14} Spearman={sp:+.3f}  {'✓方向正确' if sp < 0 else '✗方向相反'}")
    return base, acc, avg, official


def main():
    print("\n" + "#" * 104)
    analyze(os.path.join(DATA, "if_sim_2025_events.xlsx"),
            os.path.join(DATA, "if_sim_2025_baseline.xlsx"), 2025)
    print("\n" + "#" * 104)
    analyze(os.path.join(DATA, "if_sim_2024_events.xlsx"),
            os.path.join(DATA, "if_sim_2024_baseline.xlsx"), 2024)


if __name__ == "__main__":
    main()
