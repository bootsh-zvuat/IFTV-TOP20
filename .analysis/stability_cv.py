# -*- coding: utf-8 -*-
"""验证"稳定性反向"是否是真实信号，还是过拟合噪声。

观察（stability_test.py）：
  2025: (0,-1,0) → 中段MAE 4.79（基准4.57，略差）；(1,-1,0) → 4.79
  2024: (0,-1,0) → 中段MAE 4.86（基准5.21，改善）；(0,-2,0) → 5.36

两个年份"负 cv"都优于"正 cv"，但相对基准一好一坏。
cv 与官方名次的相关性：2025 是 +0.420（方向相反），2024 是 -0.059（几乎无关）。
→ 两年不一致，说明这是噪声而非稳定信号。

本脚本用**留一交叉验证**做判决：在一年上调最优 k，在另一年测。
若跨年不成立，则稳定性信号不可用。
"""
import os
from collections import defaultdict

import openpyxl
import numpy as np

BASE = r"D:\demo\iftv"
DATA = os.path.join(BASE, "data")
EVENT_LEVEL_WEIGHTS = {
    "Major": 1.0, "超级精英赛": 0.8, "精英赛": 0.75,
    "大型赛事": 0.65, "中型赛事": 0.5, "次级赛事": 0.45, "不纳入": 0.0,
}
PLACEMENT_ADJUST = {"冠军": 2, "亚军": 0.5, "四强": 0, "八强": -0.5, "小组淘汰": -1}
RATING_TIERS = [(1.3, 2), (1.2, 1), (1.1, 0.5)]
RATING_BELOW = [(1.05, -0.5), (1.0, -1)]


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def load_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if r and any(v is not None for v in r)]


def rating_bonus(r):
    b = 0.0
    for th, bo in RATING_TIERS:
        if r >= th:
            b = max(b, bo)
    for th, bo in RATING_BELOW:
        if r < th:
            b = min(b, bo)
    return b


def spearman(a, b):
    a, b = np.asarray(a, float), np.asarray(b, float)
    if len(a) < 2:
        return float("nan")
    ra = np.argsort(np.argsort(a)).astype(float)
    rb = np.argsort(np.argsort(b)).astype(float)
    ra -= ra.mean(); rb -= rb.mean()
    d = np.sqrt((ra ** 2).sum() * (rb ** 2).sum())
    return float((ra * rb).sum() / d) if d else float("nan")


def build(year, evf, blf):
    rows = load_rows(os.path.join(DATA, evf))
    official = {str(r.get("选手名")).strip(): int(num(r.get("名次")))
                for r in load_rows(os.path.join(DATA, blf))}
    acc = defaultdict(lambda: defaultdict(float))
    per_ev = defaultdict(list)
    for r in rows:
        nm = str(r.get("选手名") or "").strip()
        if not nm:
            continue
        mp = num(r.get("地图数"))
        if mp and mp < 3:
            continue
        lvl = str(r.get("赛事级别") or "").strip()
        w = num(r.get("赛事权重")) or EVENT_LEVEL_WEIGHTS.get(lvl, 0)
        if w == 0:
            continue
        a = acc[nm]
        a["mp"] += mp
        rt = num(r.get("Rating"))
        if rt > 0:
            per_ev[nm].append((rt, w))
        honor = str(r.get("荣誉类型") or "无").strip()
        rk = int(num(r.get("荣誉顺位")))
        hb = 10 if honor == "MVP" else ({1: 8, 2: 6, 3: 4}.get(rk, 0) if honor == "EVP" else 0)
        a["honor"] += hb * w
        a["rating"] += rating_bonus(rt) * w
        a["place"] += PLACEMENT_ADJUST.get(str(r.get("队伍名次") or "").strip(), 0) * w

    players = [nm for nm in acc if acc[nm]["mp"] >= 20]
    out = {}
    for nm in players:
        recs = per_ev[nm]
        if len(recs) >= 2:
            ws_ = sum(w for _, w in recs)
            mu = sum(v * w for v, w in recs) / ws_
            var = sum(w * (v - mu) ** 2 for v, w in recs) / ws_
            cv = (var ** 0.5) / mu if mu else 0.0
        else:
            cv = 0.0
        out[nm] = dict(total=acc[nm]["honor"] + acc[nm]["rating"] + acc[nm]["place"],
                       cv=cv)
    return official, out


def mid_mae(official, f, players, scores):
    pred = {nm: i + 1 for i, (nm, _) in enumerate(sorted(scores.items(), key=lambda kv: -kv[1]))}
    seg = [nm for nm in official if 7 <= official[nm] <= 20 and nm in pred]
    if not seg:
        return None
    return float(np.mean([abs(pred[nm] - official[nm]) for nm in seg])), pred


def main():
    data = {}
    for year, evf, blf in [
        (2025, "if_sim_2025_events.xlsx", "if_sim_2025_baseline.xlsx"),
        (2024, "if_sim_2024_events.xlsx", "if_sim_2024_baseline.xlsx"),
    ]:
        data[year] = build(year, evf, blf)

    print("=" * 100)
    print("跨年验证：在一年上选最优 k_cv，在另一年测试")
    print("=" * 100)

    KS = [-3, -2, -1, -0.5, 0, 0.5, 1, 2, 3]
    table = {}
    for year, (official, f) in data.items():
        players = list(f.keys())
        v = np.array([f[nm]["cv"] for nm in players], float)
        mu, sd = v.mean(), v.std()
        zcv = {nm: (f[nm]["cv"] - mu) / sd if sd > 1e-9 else 0.0 for nm in players}
        row = {}
        for k in KS:
            sc = {nm: f[nm]["total"] + k * zcv[nm] for nm in players}
            r = mid_mae(official, f, players, sc)
            row[k] = r[0] if r else float("nan")
        table[year] = row
        best_k = min(row, key=lambda k: row[k])
        print(f"\n  {year} 年:")
        for k in KS:
            mark = " ←最优" if k == best_k else ""
            print(f"    k_cv={k:>+5}: 中段MAE={row[k]:5.2f}{mark}")
        print(f"    最优 k={best_k} (MAE={row[best_k]:.2f})  基准 k=0 (MAE={row[0]:.2f})")

    print("\n" + "=" * 100)
    print("跨年判决：")
    print("=" * 100)
    for src, dst in [(2025, 2024), (2024, 2025)]:
        bk = min(table[src], key=lambda k: table[src][k])
        gain_in = table[dst][bk] - table[dst][0]
        verdict = "✓ 迁移成功" if gain_in < 0 else "✗ 迁移失败（反而更差）"
        print(f"  在 {src} 选 k={bk:>+5} → 应用到 {dst}: "
              f"中段MAE {table[dst][0]:.2f} → {table[dst][bk]:.2f} ({gain_in:+.2f}) {verdict}")

    print("\n  结论依据：若两年最优 k 同号且跨年都改善，才是真信号。")
    b25 = min(table[2025], key=lambda k: table[2025][k])
    b24 = min(table[2024], key=lambda k: table[2024][k])
    print(f"  2025 最优 k={b25:>+5}；2024 最优 k={b24:>+5}")
    print(f"  同号: {'是' if b25 * b24 > 0 else '否'}")


if __name__ == "__main__":
    main()
