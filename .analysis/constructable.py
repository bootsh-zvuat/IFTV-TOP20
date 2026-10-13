# -*- coding: utf-8 -*-
"""验证：能否用已有数据自行构造"对手强度"？

关键数据：if_sim_*_events.xlsx 每行有 (赛事名, 队伍, 队伍名次, 选手Rating)。
若已知某队伍在某赛事的名次，且知道同赛事其他队伍的名次，
可以用"同赛事其他队伍的平均实力"近似对手强度。

但更直接的问题：中段误差是否真能由"赛事内对手强度"解释？
HLTV 的 vs top-N 是"对阵 top-N 战队时的 rating"，
而我们有的是"每赛事每选手的 rating + 该赛事所有队伍的名次"。

可构造的近似信号：
  A. 赛事强度 = 该赛事参赛队的平均名次分布（有多少强队）
  B. 选手面对的对手强度 ≈ 由其队伍名次推断（走得远 = 后期遇到强队）
  C. 选手 rating 相对该赛事的分布位置（同赛事内 z-score）

其中 C 最有希望：它把"绝对 rating"转成"相对同侪的表现"，
正是官方中段在意的东西（在强赛事里 rating 1.15 比在弱赛事里 1.25 更值钱）。

本脚本检验 C 及其他可构造信号对中段的效果。
"""
import os
from collections import defaultdict
import openpyxl
import numpy as np

BASE = r"D:\demo\iftv"
DATA = os.path.join(BASE, "data")
EVENT_LEVEL_WEIGHTS = {
    "Major": 1.0, "超级精英赛": 0.8, "精英赛": 0.75,
    "大型赛事": 0.65, "中型赛事": 0.5, "次级赛事": 0.45, "不纳入": 0.0,
}
PLACEMENT_ADJUST = {"冠军": 2, "亚军": 0.5, "四强": 0, "八强": -0.5, "小组淘汰": -1}
RATING_TIERS = [(1.3, 2), (1.2, 1), (1.1, 0.5)]
RATING_BELOW = [(1.05, -0.5), (1.0, -1)]
MIN_EVENT_MAPS = 3


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def load_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if r and any(v is not None for v in r)]


def rating_bonus(r):
    b = 0.0
    for th, bo in RATING_TIERS:
        if r >= th:
            b = max(b, bo)
    for th, bo in RATING_BELOW:
        if r < th:
            b = min(b, bo)
    return b


def spearman(a, b):
    a, b = np.asarray(a, float), np.asarray(b, float)
    ra = np.argsort(np.argsort(a)).astype(float)
    rb = np.argsort(np.argsort(b)).astype(float)
    ra -= ra.mean(); rb -= rb.mean()
    d = np.sqrt((ra ** 2).sum() * (rb ** 2).sum())
    return float((ra * rb).sum() / d) if d else float("nan")


def main():
    for year, evf, blf in [
        (2025, "if_sim_2025_events.xlsx", "if_sim_2025_baseline.xlsx"),
        (2024, "if_sim_2024_events.xlsx", "if_sim_2024_baseline.xlsx"),
    ]:
        rows = load_rows(os.path.join(DATA, evf))
        official = {str(r.get("选手名")).strip(): int(num(r.get("名次")))
                    for r in load_rows(os.path.join(DATA, blf))}

        # 按赛事分组，算赛事内 rating 分布
        by_event = defaultdict(list)
        valid = []
        for r in rows:
            nm = str(r.get("选手名") or "").strip()
            if not nm:
                continue
            mp = num(r.get("地图数"))
            if mp and mp < MIN_EVENT_MAPS:
                continue
            lvl = str(r.get("赛事级别") or "").strip()
            w = num(r.get("赛事权重")) or EVENT_LEVEL_WEIGHTS.get(lvl, 0)
            if w == 0:
                continue
            rec = dict(nm=nm, mp=mp, lvl=lvl, w=w,
                       rating=num(r.get("Rating")),
                       place=str(r.get("队伍名次") or "").strip(),
                       team=str(r.get("队伍") or "").strip(),
                       honor=str(r.get("荣誉类型") or "无").strip(),
                       rank=int(num(r.get("荣誉顺位"))),
                       event=str(r.get("赛事名称") or "").strip())
            valid.append(rec)
            by_event[rec["event"]].append(rec)

        # 赛事内 rating 均值和标准差（用于相对 z-score）
        ev_stats = {}
        for evname, recs in by_event.items():
            rs = [x["rating"] for x in recs if x["rating"] > 0]
            if len(rs) >= 5:
                ev_stats[evname] = (float(np.mean(rs)), float(np.std(rs)) or 1e-9)
            else:
                ev_stats[evname] = (None, None)

        # 赛事强度 = 参赛队的名次分布（冠亚四强比例）
        ev_strength = {}
        for evname, recs in by_event.items():
            teams = {}
            for x in recs:
                if x["team"]:
                    teams[x["team"]] = x["place"]
            if not teams:
                ev_strength[evname] = 0.5
                continue
            vals = {"冠军": 1.0, "亚军": 0.8, "四强": 0.6, "八强": 0.4, "小组淘汰": 0.2}
            ev_strength[evname] = float(np.mean([vals.get(v, 0.2) for v in teams.values()]))

        acc = defaultdict(lambda: defaultdict(float))
        zratings = defaultdict(list)
        for x in valid:
            a = acc[x["nm"]]
            a["mp"] += x["mp"]
            a["rating"] += rating_bonus(x["rating"]) * x["w"]
            a["place"] += PLACEMENT_ADJUST.get(x["place"], 0) * x["w"]
            hb = 10 if x["honor"] == "MVP" else (
                {1: 8, 2: 6, 3: 4}.get(x["rank"], 0) if x["honor"] == "EVP" else 0)
            a["honor"] += hb * x["w"]
            mu, sd = ev_stats.get(x["event"], (None, None))
            if mu is not None and x["rating"] > 0:
                z = (x["rating"] - mu) / sd
                zratings[x["nm"]].append((z, x["w"]))

        players = [nm for nm in acc if acc[nm]["mp"] >= 20]
        base = {nm: acc[nm]["honor"] + acc[nm]["rating"] + acc[nm]["place"] for nm in players}
        # 赛事内相对 rating（加权平均 z）
        relz = {}
        for nm in players:
            zs = zratings.get(nm) or []
            if zs:
                tot_w = sum(w for _, w in zs)
                relz[nm] = sum(z * w for z, w in zs) / tot_w if tot_w else 0.0
            else:
                relz[nm] = 0.0

        def z(d):
            v = np.array([d[nm] for nm in players], float)
            mu, sd = v.mean(), v.std()
            return {nm: (d[nm] - mu) / sd if sd > 1e-9 else 0.0 for nm in players}

        zrel = z(relz)
        zbase = z(base)

        def ev(label, scores):
            rk = sorted(scores.items(), key=lambda kv: -kv[1])
            pred = {nm: i + 1 for i, (nm, _) in enumerate(rk)}
            mid = [nm for nm in official if 7 <= official[nm] <= 20 and nm in pred]
            head = [nm for nm in official if official[nm] <= 6 and nm in pred]
            top20 = len({nm for nm in official if nm in list(pred)[:20]})
            if not mid:
                print(f"  {label:<38} 无数据")
                return
            mae = float(np.mean([abs(pred[nm] - official[nm]) for nm in mid]))
            sp = spearman([official[nm] for nm in mid], [pred[nm] for nm in mid])
            hmae = float(np.mean([abs(pred[nm] - official[nm]) for nm in head])) if head else float("nan")
            print(f"  {label:<38} Top20={top20:>2}/20 头部MAE={hmae:4.2f} "
                  f"中段MAE={mae:5.2f} 中段sp={sp:+.3f}")

        print("=" * 100)
        print(f"{year} 年：可自行构造的对手强度信号")
        print("=" * 100)
        ev("现行 v1.2.0", base)
        ev("z(base) + z(赛事内相对rating)×k, k=1", {nm: zbase[nm] + 1 * zrel[nm] for nm in players})
        ev("z(base) + z(赛事内相对rating)×k, k=3", {nm: zbase[nm] + 3 * zrel[nm] for nm in players})
        ev("z(赛事内相对rating) 单项", {nm: zrel[nm] for nm in players})
        strength_proxy = {nm: ev_strength_proxy(valid, ev_strength, nm) for nm in players}
        ev("z(base) + z(赛事强度)", {nm: zbase[nm] + z(strength_proxy)[nm] for nm in players})
        print(f"\n  中段内 赛事内相对rating 与官方名次 Spearman = "
              f"{spearman([official[n] for n in players if n in official and 7 <= official[n] <= 20], [relz[n] for n in players if n in official and 7 <= official[n] <= 20]):+.3f}")
        print()


def ev_strength_proxy(valid, ev_strength, nm):
    """选手加权平均面对的赛事强度"""
    tot = 0.0
    wsum = 0.0
    for x in valid:
        if x["nm"] != nm:
            continue
        s = ev_strength.get(x["event"], 0.5)
        tot += s * x["w"]
        wsum += x["w"]
    return tot / wsum if wsum else 0.0


if __name__ == "__main__":
    main()
