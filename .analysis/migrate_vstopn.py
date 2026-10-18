# -*- coding: utf-8 -*-
"""给已有赛事数据补上 vsTopN 列（无需重新采集）。

背景：
  gen_if_data_2025.py 已支持输出 vsTopN 列，但完整重跑需要重新采集
  26 个赛事 × 数百场比赛（数小时 API 调用）。而 vsTopN 是**选手年度聚合值**，
  只依赖 data/hltv_opponent_<year>.xlsx，与赛事采集无关。

  因此直接给现有 xlsx 追加 5 列即可，按选手名 join。

注意：这是**数据迁移**，不改变任何已有列的值（先备份）。
"""
import os
import shutil

import openpyxl
from openpyxl.styles import Font, PatternFill

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 项目根 = .analysis/ 的上一级
DATA = os.path.join(BASE, "data")
KEYS = ["vsTop5", "vsTop10", "vsTop20", "vsTop30", "vsTop50"]

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def norm_name(s):
    import re
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", (s or "").lower())


def load_opp(year):
    p = os.path.join(DATA, f"hltv_opponent_{year}.xlsx")
    if not os.path.exists(p):
        return {}
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = {}
    for r in rows[1:]:
        if not r or all(v is None for v in r):
            continue
        d = dict(zip(hdr, r))
        nm = norm_name(str(d.get("选手") or ""))
        if nm:
            out[nm] = {k: d.get(k) for k in KEYS}
    return out


def migrate(year):
    path = os.path.join(DATA, f"if_sim_{year}_events.xlsx")
    if not os.path.exists(path):
        print(f"  ! 缺少 {path}")
        return False

    wb = openpyxl.load_workbook(path)
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    print(f"  {year}: 原列数={len(hdr)}")

    if any(h and "vsTop" in str(h) for h in hdr):
        print(f"  {year}: 已含 vsTop 列，跳过")
        wb.close()
        return True

    # 备份
    bak = path + ".bak"
    shutil.copy2(path, bak)
    print(f"  已备份 -> {os.path.basename(bak)}")

    opp = load_opp(year)
    print(f"  HLTV 对手强度载入 {len(opp)} 人")

    # 追加表头
    start_col = len(hdr) + 1
    for i, k in enumerate(KEYS):
        c = ws.cell(row=1, column=start_col + i, value=k)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT

    # 找选手名列
    try:
        name_idx = [str(h).strip() for h in hdr].index("选手名") + 1
    except ValueError:
        print("  ! 找不到「选手名」列，终止")
        wb.close()
        return False

    filled = 0
    total = 0
    for row in range(2, ws.max_row + 1):
        nm = ws.cell(row=row, column=name_idx).value
        if nm is None:
            continue
        total += 1
        o = opp.get(norm_name(str(nm)))
        for i, k in enumerate(KEYS):
            if o and o.get(k) not in (None, ""):
                ws.cell(row=row, column=start_col + i, value=o[k])
                filled += 1
    wb.save(path)
    wb.close()
    print(f"  完成：{total} 行，{filled} 个 vsTop 值写入（{filled//len(KEYS)} 行有数据）")
    return True


def main():
    print("=" * 80)
    print("给赛事数据补 vsTopN 列（按选手名 join，先备份）")
    print("=" * 80)
    for y in (2025, 2024):
        migrate(y)
        print()
    print("提示：这是年度聚合值，同一选手全年各赛事行相同（符合 HLTV 口径）。")
    print("      .bak 备份可删；确认无误后建议清理。")


if __name__ == "__main__":
    main()
