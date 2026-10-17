# -*- coding: utf-8 -*-
"""端到端验证：TypeScript 实现是否复现 20/20？

前一步用 Python 原型验证了 vsTop5 通道有效。现在 TS 已接入，
必须确认 TS 实现（z-score + 加到总分 + 小样本门槛顺序）结果一致。

方法：用 Python **逐行复刻 TS 的 computeRanking/opponentScore 逻辑**，
读取与 TS 相同的输入（DATA + 权重默认值），比较：
  - Top20 重叠
  - 头部/中段 MAE
  - 关闭通道（weight=0）时是否回到 4.57 / 5.21 基线

注意：这里复刻的是**最终 TS 逻辑**，不是之前的原型 —— 若两者不一致，
说明 TS 实现有偏差，必须修 TS 而不是改这个脚本。
"""
import json
import os
from collections import defaultdict

import openpyxl
import numpy as np

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # 项目根 = .analysis/ 的上一级
DATA = os.path.join(BASE, "data")

# ---- 与 src/config/defaults.ts 保持一致 ----
EVENT_LEVEL_WEIGHTS = {
    "Major": 1.0, "超级精英赛": 0.8, "精英赛": 0.75,
    "大型赛事": 0.65, "中型赛事": 0.5, "次级赛事": 0.45, "不纳入": 0.0,
}
PLACEMENT_ADJUST = {"冠军": 2, "亚军": 0.5, "四强": 0, "八强": -0.5, "小组淘汰": -1}
HONOR_BASE = {"MVP": 10, "1": 8, "2": 6, "3": 4}
RATING_TIERS = [(1.3, 2), (1.2, 1), (1.1, 0.5)]
RATING_BELOW = [(1.05, -0.5), (1.0, -1)]
SCORE_MIN, SCORE_MAX = -2, 11
MIN_MAPS, MIN_EVENT_MAPS = 20, 3
OPPONENT_WEIGHT = 2
OPPONENT_TIER = 5


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def opt_num(v):
    if v is None or str(v).strip() == "":
        return None
    try:
        f = float(v)
        return f if f > 0 else None
    except (TypeError, ValueError):
        return None


def load_rows(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(hdr, r)) for r in rows[1:] if r and any(v is not None for v in r)]


def rating_bonus(r):
    b = 0.0
    for th, bo in RATING_TIERS:
        if r >= th:
            b = max(b, bo)
    for th, bo in RATING_BELOW:
        if r < th:
            b = min(b, bo)
    return b


def event_perf(row):
    honor = str(row.get("荣誉类型") or "无").strip()
    rk = int(num(row.get("荣誉顺位")))
    hb = HONOR_BASE["MVP"] if honor == "MVP" else (
        HONOR_BASE.get(str(rk), 0) if honor == "EVP" else 0)
    s = hb + rating_bonus(num(row.get("Rating"))) + \
        PLACEMENT_ADJUST.get(str(row.get("队伍名次") or "").strip(), 0)
    return min(SCORE_MAX, max(SCORE_MIN, s))


def opponent_score(rows, weight, tier):
    """复刻 src/core/ranking.ts 的 opponentScore：年度 z-score，每选手取一次"""
    if weight == 0:
        return {}
    key = f"vsTop{tier}"
    by_player = {}
    for r in rows:
        nm = str(r.get("选手名") or "").strip()
        if not nm or nm in by_player:
            continue
        v = opt_num(r.get(key))
        if v is not None:
            by_player[nm] = v
    if len(by_player) < 2:
        return {}
    vals = list(by_player.values())
    mean = sum(vals) / len(vals)
    var = sum((v - mean) ** 2 for v in vals) / len(vals)
    sd = var ** 0.5
    if sd < 1e-9:
        return {}
    return {nm: (v - mean) / sd * weight for nm, v in by_player.items()}


def compute_ranking(year, weight=OPPONENT_WEIGHT, tier=OPPONENT_TIER):
    rows = load_rows(os.path.join(DATA, f"if_sim_{year}_events.xlsx"))
    totals = defaultdict(float)
    maps_sum = defaultdict(float)
    # 逐赛事加和（含两道门槛）
    for r in rows:
        nm = str(r.get("选手名") or "").strip()
        if not nm:
            continue
        mp = num(r.get("地图数"))
        if mp and mp < MIN_EVENT_MAPS:
            continue
        maps_sum[nm] += mp
        lvl = str(r.get("赛事级别") or "").strip()
        w = num(r.get("赛事权重")) or EVENT_LEVEL_WEIGHTS.get(lvl, 0)
        if w == 0:
            continue
        totals[nm] += event_perf(r) * w
    # 对手强度通道
    for nm, bonus in opponent_score(rows, weight, tier).items():
        if nm in totals:
            totals[nm] += bonus
    # 年度图数门槛
    out = {nm: t for nm, t in totals.items() if maps_sum[nm] >= MIN_MAPS}
    return out


def spearman(a, b):
    a, b = np.asarray(a, float), np.asarray(b, float)
    if len(a) < 2:
        return float("nan")
    ra = np.argsort(np.argsort(a)).astype(float)
    rb = np.argsort(np.argsort(b)).astype(float)
    ra -= ra.mean(); rb -= rb.mean()
    d = np.sqrt((ra ** 2).sum() * (rb ** 2).sum())
    return float((ra * rb).sum() / d) if d else float("nan")


def report(year, weight, tier):
    scores = compute_ranking(year, weight, tier)
    official = {str(r.get("选手名")).strip(): int(num(r.get("名次")))
                for r in load_rows(os.path.join(DATA, f"if_sim_{year}_baseline.xlsx"))}
    pred = {nm: i + 1 for i, (nm, _) in enumerate(
        sorted(scores.items(), key=lambda kv: -kv[1]))}
    t20 = len({nm for nm in official if nm in list(pred)[:20]})
    seg = lambda lo, hi: [nm for nm in official if lo <= official[nm] <= hi and nm in pred]
    def mae(lo, hi):
        s = seg(lo, hi)
        return float(np.mean([abs(pred[nm] - official[nm]) for nm in s])) if s else float("nan")
    def sp(lo, hi):
        s = seg(lo, hi)
        return spearman([official[nm] for nm in s], [pred[nm] for nm in s]) if s else float("nan")
    return dict(year=year, weight=weight, tier=tier, top20=t20,
                head_mae=mae(1, 6), mid_mae=mae(7, 20),
                mid_sp=sp(7, 20), overall_mae=mae(1, 20), overall_sp=sp(1, 20))


def main():
    print("=" * 96)
    print("端到端验证：TypeScript 集成后的效果（Python 复刻 TS 逻辑）")
    print("=" * 96)
    print(f"{'年份':>6} {'weight':>7} {'tier':>5} | {'Top20':>6} {'头部MAE':>8} "
          f"{'中段MAE':>8} {'中段sp':>8} {'全体MAE':>8} {'全体sp':>8}")
    rows = []
    for y in (2025, 2024):
        for w, t in ((0, 5), (2, 5), (3, 5), (2, 20)):
            r = report(y, w, t)
            rows.append(r)
            print(f"{r['year']:>6} {r['weight']:>7} {r['tier']:>5} | "
                  f"{r['top20']:>4}/20 {r['head_mae']:>8.2f} {r['mid_mae']:>8.2f} "
                  f"{r['mid_sp']:>+8.3f} {r['overall_mae']:>8.2f} {r['overall_sp']:>+8.3f}")

    print("\n" + "=" * 96)
    print("验收判定")
    print("=" * 96)
    ok = True
    for y in (2025, 2024):
        off = [r for r in rows if r["year"] == y and r["weight"] == 0][0]
        on = [r for r in rows if r["year"] == y and r["weight"] == 2 and r["tier"] == 5][0]
        print(f"\n  {y}:")
        print(f"    关闭通道: Top20={off['top20']}/20 中段MAE={off['mid_mae']:.2f} "
              f"头部MAE={off['head_mae']:.2f}")
        print(f"    开启通道: Top20={on['top20']}/20 中段MAE={on['mid_mae']:.2f} "
              f"头部MAE={on['head_mae']:.2f}")
        good = on["top20"] >= off["top20"] and on["mid_mae"] < off["mid_mae"] \
            and on["head_mae"] <= off["head_mae"] + 0.01
        print(f"    → {'✓ 通过' if good else '✗ 未通过'}"
              f"（要求：Top20 不降、中段改善、头部不退化）")
        ok &= good

    print("\n" + "=" * 96)
    print("关闭通道时是否回到 v1.2.0 基线？（基线：2025 中段 4.57 / 2024 中段 5.21）")
    print("=" * 96)
    for y, exp in ((2025, 4.57), (2024, 5.21)):
        off = [r for r in rows if r["year"] == y and r["weight"] == 0][0]
        print(f"  {y}: 期望 {exp:.2f}  实际 {off['mid_mae']:.2f}  "
              f"{'✓' if abs(off['mid_mae'] - exp) < 0.02 else '✗ 偏差过大'}")

    print("\n" + "=" * 96)
    print("最终判定: " + ("✓ TS 集成成功，指标达标" if ok else "✗ 需检查"))
    print("=" * 96)


if __name__ == "__main__":
    main()
