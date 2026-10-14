# -*- coding: utf-8 -*-
"""
HLTV 对手强度采集器（通过 Edge CDP，零第三方依赖）。

原理：
  真实 Edge 窗口能正常通过 Cloudflare 的 JS 挑战（实测 29s），
  再用 Edge 内置的 CDP 接口读取渲染后的 DOM。

  这不是"绕过"保护 —— 浏览器自己完成了挑战，脚本只是自动点开页面并读结果。
  相比人工导出，好处是可批量、可复现；相比伪造请求，好处是行为与真人一致。

关键限制（务必知晓）：
  · 需要**有头** Edge（headless 会被 CF 挡住，已实测）
  · 首次运行会弹出浏览器窗口，属正常现象
  · 每页约 5-30s，20 人/年约 5-10 分钟

用法：
  1. 启动 Edge（独立 profile，避免影响日常浏览）：
     "C:\\Program Files (x86)\\Microsoft\\Edge\\Application\\msedge.exe" ^
       --remote-debugging-port=9411 ^
       --user-data-dir="D:\\demo\\iftv\\.analysis\\edge_hltv_profile" ^
       --disable-blink-features=AutomationControlled about:blank

  2. 运行：
     .venv\\Scripts\\python.exe fetch_hltv_opponents.py --year 2025 --out data\\hltv_opponent_2025.xlsx
"""
import argparse
import os
import re
import sys
import time

# 项目根 = src/ 的上一级；cdp_client.py 在根目录的 .analysis/ 下
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, ".analysis"))
from cdp_client import WS, cdp_get  # noqa: E402

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    print("缺少 openpyxl"); sys.exit(1)


# ---------- CDP 会话 ----------
class Session:
    def __init__(self, timeout=25):
        u = None
        for t in cdp_get("/json/list"):
            if t.get("type") == "page":
                u = t.get("webSocketDebuggerUrl")
                break
        if not u:
            raise RuntimeError(
                "未找到 Edge 页面 target。请先用 --remote-debugging-port=9411 启动 Edge。")
        self.ws = WS(u, timeout=timeout)
        self.mid = 1
        self.ws.call("Page.enable", msg_id=self.mid); self.mid += 1

    def ev(self, expr, timeout=30):
        r = self.ws.call("Runtime.evaluate",
                         {"expression": expr, "returnByValue": True}, msg_id=self.mid)
        self.mid += 1
        res = r.get("result", {})
        if "exceptionDetails" in res:
            return None
        return res.get("result", {}).get("value")

    def goto(self, url, wait_for, max_wait=90):
        """导航并等待特征文本出现（自动穿过 CF 挑战页）"""
        self.ws.call("Page.navigate", {"url": url}, msg_id=self.mid)
        self.mid += 1
        t0 = time.time()
        while time.time() - t0 < max_wait:
            time.sleep(1)
            txt = self.ev("document.body ? document.body.innerText : ''") or ""
            if wait_for(txt):
                return txt
        return None


# ---------- 数据提取（基于实测 DOM：div.rating-breakdown） ----------
EXTRACT_JS = r"""
(() => {
  const out = {vs: {}, maps: {}, name: '', summary: {}};

  // vs Top-N：实测结构
  //   <div class="rating-breakdown">
  //     <div class="rating-value">1.33</div>
  //     <div class="rating-description">vs top 5 opponents</div>
  //     <div class="rating-maps">(63 maps)</div>
  //   </div>
  document.querySelectorAll('.rating-breakdown').forEach(el => {
    const desc = (el.querySelector('.rating-description')||{}).innerText || '';
    const val  = (el.querySelector('.rating-value')||{}).innerText || '';
    const maps = (el.querySelector('.rating-maps')||{}).innerText || '';
    const m = desc.match(/vs\s*top\s*(\d+)/i);
    if (m) {
      out.vs[m[1]] = parseFloat(val);
      const mm = maps.match(/(\d+)/);
      if (mm) out.maps[m[1]] = parseInt(mm[1]);
    }
  });

  // 选手名：h1 不可靠（实测被侧栏 RECENT ACTIVITY 占据），
  // 从 <title> 提取： "Mathieu 'ZywOo' Herbaut Counter-Strike Statistics | HLTV.org"
  // → 取单引号内的昵称；无引号则取 title 首段
  const t = document.title || '';
  const q = t.match(/'([^']+)'/);
  out.name = q ? q[1] : t.split('|')[0].trim();

  // 顶部汇总：实测结构为 <div class="stats-row">标签\n值</div>
  // 例： "K/D Ratio\n1.42"、"Maps played\n200"、"Impact rating\n1.37"
  document.querySelectorAll('.stats-row').forEach(el => {
    const parts = (el.innerText || '').split('\n').map(s => s.trim()).filter(Boolean);
    if (parts.length === 2) out.summary[parts[0]] = parts[1];
  });

  // 注意：此页面**没有**独立的"整体 Rating"。
  // 实测 .rating-value 恰好只有 5 个，且与 5 个 vs-Top-N 分档一一对应
  // （ZywOo: 1.33/1.33/1.36/1.34/1.34 = vs5/vs10/vs20/vs30/vs50）。
  // 早期版本误把首个 .rating-value 当作整体 Rating，导致 Rating 列与 vsTop5
  // 完全重复（r=1.000），进而使验证结果虚假改善 —— 已修正为不再输出该列。
  return out;
})()
"""


def fetch_player(sess, player_id, slug, year):
    url = (f"https://www.hltv.org/stats/players/{player_id}/{slug}"
           f"?startDate={year}-01-01&endDate={year}-12-31")
    txt = sess.goto(url, lambda t: "rating-breakdown" in t or "vs top 5" in t.lower(),
                    max_wait=90)
    if not txt:
        return None
    data = sess.ev(EXTRACT_JS)
    if not data:
        return None
    vs = {int(k): v for k, v in (data.get("vs") or {}).items() if v is not None}
    maps = {int(k): v for k, v in (data.get("maps") or {}).items()}
    if not vs:
        return None
    summ = data.get("summary") or {}
    return dict(
        name=data.get("name") or "",
        vs=vs, vs_maps=maps,
        # 不输出"整体 Rating"：该页面没有此项，硬凑会与 vsTop5 完全重复。
        # 需要整体 rating 时请用赛事数据里的 Rating（5E/完美口径）或 vs50。
        kd_ratio=_num(summ.get("K/D Ratio")),
        impact=_num(summ.get("Impact rating")),
        adr=_num(summ.get("Damage / Round")),
        kpr=_num(summ.get("Kills / round")),
        maps_played=_num(summ.get("Maps played")) or maps.get(50),
        url=url,
    )


def _num(s):
    if s is None:
        return None
    m = re.search(r"-?\d+(?:\.\d+)?", str(s).replace(",", ""))
    return float(m.group(0)) if m else None


# ---------- 选手清单 ----------
def load_top20(year):
    """从 baseline xlsx 读该年 Top20 的 (选手名, 选手ID)"""
    path = os.path.join(ROOT, "data", f"if_sim_{year}_baseline.xlsx")
    if not os.path.exists(path):
        return []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or all(v is None for v in r):
            continue
        d = dict(zip(hdr, r))
        name = str(d.get("选手名") or "").strip()
        pid = d.get("选手ID")
        if name:
            out.append((name, pid))
    return out


def main():
    ap = argparse.ArgumentParser(description="通过 Edge CDP 采集 HLTV 对手强度数据")
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--limit", type=int, default=20)
    ap.add_argument("--list", help="可选：选手清单文件（每行 选手名,选手ID）")
    args = ap.parse_args()

    players = load_top20(args.year)
    if args.list:
        players = []
        with open(args.list, encoding="utf-8") as f:
            for ln in f:
                ln = ln.strip()
                if not ln:
                    continue
                parts = ln.split(",")
                players.append((parts[0], parts[1] if len(parts) > 1 else None))
    if not players:
        print(f"未找到 {args.year} 年选手清单（需要 data/if_sim_{args.year}_baseline.xlsx）")
        return 1
    players = players[:args.limit]
    print(f"待采集 {len(players)} 人（{args.year} 年）\n")

    sess = Session()
    rows, failed = [], []
    for i, (name, pid) in enumerate(players, 1):
        if not pid:
            print(f"  [{i}/{len(players)}] {name}: 缺选手ID，跳过")
            failed.append(name)
            continue
        slug = name
        try:
            d = fetch_player(sess, pid, slug, args.year)
        except Exception as e:
            print(f"  [{i}/{len(players)}] {name}: 异常 {str(e)[:60]}")
            d = None
        if not d or not d["vs"]:
            print(f"  [{i}/{len(players)}] {name}: ❌ 未取到 vs Top-N")
            failed.append(name)
            continue
        print(f"  [{i}/{len(players)}] {d['name']:<14} "
              f"vs5={d['vs'].get(5)} vs10={d['vs'].get(10)} vs20={d['vs'].get(20)} "
              f"maps={d['maps_played']}")
        rows.append([args.year, d["name"], d["maps_played"],
                     d["kd_ratio"], d["impact"], d["adr"], d["kpr"],
                     d["vs"].get(5), d["vs"].get(10), d["vs"].get(20),
                     d["vs"].get(30), d["vs"].get(50),
                     d["vs_maps"].get(5), d["vs_maps"].get(20), d["url"]])

    if not rows:
        print("\n没有任何数据被采集，终止。")
        return 1

    header = ["年份", "选手", "地图数", "KD比值", "Impact", "ADR", "KPR",
              "vsTop5", "vsTop10", "vsTop20", "vsTop30", "vsTop50",
              "vsTop5图数", "vsTop20图数", "来源URL"]
    wb = Workbook(); ws = wb.active; ws.title = "opponent"
    fill = PatternFill("solid", fgColor="305496")
    font = Font(bold=True, color="FFFFFF")
    for c, h in enumerate(header, 1):
        cell = ws.cell(row=1, column=c, value=h); cell.fill = fill; cell.font = font
    for i, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=i, column=c, value=v)
    wb.save(args.out)

    print(f"\n已写出: {args.out}（{len(rows)} 行）")
    if failed:
        print(f"! {len(failed)} 人未采集: {', '.join(failed)}")
        print("  通常原因：选手ID 不匹配 / 页面 slug 不对 / 该年无数据")
    return 0


if __name__ == "__main__":
    sys.exit(main())
